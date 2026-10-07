"""RequirementsLoader — loads requirements and principles from an xlsx file.

Column names are always read from the injected config dict, never hardcoded.
"""

import warnings
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

_OPTIONAL_COLUMNS = {"plateau", "apr_ref", "ad_ref"}


def _get_column_index(sheet: Worksheet, column_name: str) -> int | None:
    """Return 1-based column index for a header name, or None if absent."""
    for cell in sheet[1]:
        if cell.value and str(cell.value).strip() == column_name:
            return int(cell.column)
    return None


def _sheet_headers(sheet: Worksheet) -> list[str]:
    """Return non-empty header values from the first row of a sheet."""
    return [str(c.value).strip() for c in sheet[1] if c.value is not None]


class RequirementsLoader:
    """Loads requirements and principles from an xlsx file.

    Args:
        config: The parsed config.yaml dict (must contain 'xlsx' section).
    """

    def __init__(self, config: dict[str, Any]) -> None:
        self._cfg = config

    def load(self, xlsx_path: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
        """Load requirements and principles from xlsx_path.

        Returns:
            Tuple of (requirement_rows, principle_rows).  Each row is a dict
            keyed by the logical column names defined in config.yaml.

        Raises:
            FileNotFoundError: If xlsx_path does not exist.
            ValueError: If a required sheet or column is missing.
        """
        if not xlsx_path.exists():
            raise FileNotFoundError(f"Requirements file not found: {xlsx_path}")

        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        cfg_xlsx = self._cfg["xlsx"]
        req_sheet_name: str = cfg_xlsx["requirements_sheet"]
        pri_sheet_name: str = cfg_xlsx["principles_sheet"]

        if req_sheet_name not in wb.sheetnames:
            raise ValueError(
                f"xlsx: sheet '{req_sheet_name}' not found. "
                f"Available sheets: {wb.sheetnames}. "
                f"Update 'xlsx.requirements_sheet' in config.yaml."
            )
        if pri_sheet_name not in wb.sheetnames:
            raise ValueError(
                f"xlsx: sheet '{pri_sheet_name}' not found. "
                f"Available sheets: {wb.sheetnames}. "
                f"Update 'xlsx.principles_sheet' in config.yaml."
            )

        col_map: dict[str, str] = cfg_xlsx["columns"]
        req_sheet = wb[req_sheet_name]

        col_idx: dict[str, int | None] = {}
        for key, header in col_map.items():
            idx = _get_column_index(req_sheet, header)
            if idx is None and key in _OPTIONAL_COLUMNS:
                warnings.warn(f"xlsx: column '{header}' ({key}) not found — will be None")
            elif idx is None:
                available = _sheet_headers(req_sheet)
                raise ValueError(
                    f"xlsx: required column '{header}' not found in sheet '{req_sheet_name}'.\n"
                    f"  Available columns: {available}\n"
                    f"  Update 'xlsx.columns.{key}' in config.yaml to match."
                )
            col_idx[key] = idx

        rows: list[dict[str, Any]] = []
        for row in req_sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            r: dict[str, Any] = {}
            for key, idx in col_idx.items():
                r[key] = row[idx - 1] if idx is not None else None
            rows.append(r)

        pri_col_map: dict[str, str] = cfg_xlsx["principles_columns"]
        pri_sheet = wb[pri_sheet_name]
        pri_col_idx: dict[str, int] = {}
        for key, header in pri_col_map.items():
            idx = _get_column_index(pri_sheet, header)
            if idx is None:
                available = _sheet_headers(pri_sheet)
                raise ValueError(
                    f"xlsx: required column '{header}' not found in sheet '{pri_sheet_name}'.\n"
                    f"  Available columns: {available}\n"
                    f"  Update 'xlsx.principles_columns.{key}' in config.yaml to match."
                )
            pri_col_idx[key] = idx

        principles: list[dict[str, Any]] = []
        for row in pri_sheet.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            principles.append({key: row[idx - 1] for key, idx in pri_col_idx.items()})

        return rows, principles
